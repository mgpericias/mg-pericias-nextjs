#!/usr/bin/env python3
"""
Gera public/obra.html (Portfólio de Obras) a partir de dados/obras/DADOS_CLIENTES.xlsx.

Uso:
    python scripts/gerar-obra.py                 # usa os caminhos padrão
    python scripts/gerar-obra.py --planilha X    # outra planilha
    python scripts/gerar-obra.py --saida Y       # outro arquivo de saída

Requisitos: pip install openpyxl

Regras:
- Uma linha da planilha = uma obra. Obras com o mesmo "Nome do Prédio" viram um único card.
- Foto, avaliação Google e vídeo são lidos da primeira linha do prédio que tiver o campo preenchido.
- Fotos ficam em public/images/obras/ e são referenciadas como /images/obras/<arquivo>.
- O filtro de estágio só mostra os estágios que têm pelo menos uma obra.
"""
import argparse
import html
import os
import re
import sys
import unicodedata
from urllib.parse import quote

try:
    import openpyxl
except ImportError:
    sys.exit("Falta a biblioteca openpyxl. Instale com: pip install openpyxl")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANILHA = os.path.join(RAIZ, "dados", "obras", "DADOS_CLIENTES.xlsx")
TEMPLATE = os.path.join(RAIZ, "scripts", "obra-template.html")
SAIDA = os.path.join(RAIZ, "public", "obra.html")
PASTA_FOTOS = os.path.join(RAIZ, "public", "images", "obras")
URL_FOTOS = "/images/obras/"

TIPOS = {
    "gestao": ("gestao", '<span class="obra-tipo-badge tipo-gestao">Gestão</span>'),
    "fiscalizacao": ("fiscalizacao", '<span class="obra-tipo-badge tipo-fiscalizacao">Fiscalização</span>'),
    "gestao e fiscalizacao": ("gef", '<span class="obra-tipo-badge tipo-gef"><span class="g">Gestão</span>'
                                     '<span class="sep">/</span><span class="f">Fiscalização</span></span>'),
}
ESTAGIOS = [  # ordem dos botões do filtro
    ("contratacao", "obra em contratacao", "Obra em Contratação"),
    ("execucao", "em execucao", "Em Execução"),
    ("concluido", "concluido", "Concluído"),
]

SVG_PIN = ('<svg viewBox="0 0 24 24"><path d="M21 10c0 7-9 13-9 13S3 17 3 10a9 9 0 0 1 18 0z"/>'
           '<circle cx="12" cy="10" r="3"/></svg>')
SVG_CAL = ('<svg viewBox="0 0 24 24"><rect x="3" y="4" width="18" height="18" rx="2"/><line x1="16" y1="2" x2="16" y2="6"/>'
           '<line x1="8" y1="2" x2="8" y2="6"/><line x1="3" y1="10" x2="21" y2="10"/></svg>')
SVG_ZAP = ('<svg viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg"><path d="M17.472 14.382c-.297-.149-1.758-.867-2.03-.967-.273-.099-.471-.148-.67.15-.197.297-.767.966-.94 1.164-.173.199-.347.223-.644.075-.297-.15-1.255-.463-2.39-1.475-.883-.788-1.48-1.761-1.653-2.059-.173-.297-.018-.458.13-.606.134-.133.298-.347.446-.52.149-.174.198-.298.298-.497.099-.198.05-.371-.025-.52-.075-.149-.669-1.612-.916-2.207-.242-.579-.487-.5-.669-.51-.173-.008-.371-.01-.57-.01-.198 0-.52.074-.792.372-.272.297-1.04 1.016-1.04 2.479 0 1.462 1.065 2.875 1.213 3.074.149.198 2.096 3.2 5.077 4.487.709.306 1.262.489 1.694.625.712.227 1.36.195 1.871.118.571-.085 1.758-.719 2.006-1.413.248-.694.248-1.289.173-1.413-.074-.124-.272-.198-.57-.347m-5.421 7.403h-.004a9.87 9.87 0 01-5.031-1.378l-.361-.214-3.741.982.998-3.648-.235-.374a9.86 9.86 0 01-1.51-5.26c.001-5.45 4.436-9.884 9.888-9.884 2.64 0 5.122 1.03 6.988 2.898a9.825 9.825 0 012.893 6.994c-.003 5.45-4.437 9.884-9.885 9.884m8.413-18.297A11.815 11.815 0 0012.05 0C5.495 0 .16 5.335.157 11.892c0 2.096.547 4.142 1.588 5.945L.057 24l6.305-1.654a11.882 11.882 0 005.683 1.448h.005c6.554 0 11.89-5.335 11.893-11.893a11.821 11.821 0 00-3.48-8.413z"/></svg>')
SVG_GOOGLE = ('<svg class="google-icon" viewBox="0 0 24 24"><path d="M22.56 12.25c0-.78-.07-1.53-.2-2.25H12v4.26h5.92c-.26 1.37-1.04 2.53-2.21 3.31v2.77h3.57c2.08-1.92 3.28-4.74 3.28-8.09z" fill="#4285F4"/><path d="M12 23c2.97 0 5.46-.98 7.28-2.66l-3.57-2.77c-.98.66-2.23 1.06-3.71 1.06-2.86 0-5.29-1.93-6.16-4.53H2.18v2.84C3.99 20.53 7.7 23 12 23z" fill="#34A853"/><path d="M5.84 14.09c-.22-.66-.35-1.36-.35-2.09s.13-1.43.35-2.09V7.07H2.18C1.43 8.55 1 10.22 1 12s.43 3.45 1.18 4.93l3.66-2.84z" fill="#FBBC05"/><path d="M12 5.38c1.62 0 3.06.56 4.21 1.64l3.15-3.15C17.45 2.09 14.97 1 12 1 7.7 1 3.99 3.47 2.18 7.07l3.66 2.84c.87-2.6 3.3-4.53 6.16-4.53z" fill="#EA4335"/></svg>')
SVG_STAR = ('<svg viewBox="0 0 24 24"><polygon class="star-full" points="12 2 15.09 8.26 22 9.27 17 14.14 18.18 21.02 '
            '12 17.77 5.82 21.02 7 14.14 2 9.27 8.91 8.26 12 2"/></svg>')
SVG_EXT = '<svg viewBox="0 0 24 24"><line x1="7" y1="17" x2="17" y2="7"/><polyline points="7 7 17 7 17 17"/></svg>'
SVG_PLAY = '<svg viewBox="0 0 24 24"><path d="M8 5v14l11-7z"/></svg>'
SVG_FOTO = ('<svg viewBox="0 0 24 24"><rect x="3" y="3" width="18" height="18" rx="2"/><circle cx="8.5" cy="8.5" r="1.5"/>'
            '<path d="M21 15l-5-5L5 21"/></svg>')

MSG_SINDICO = ("Olá {nome}! Eu vim pelo Portfólio da MG Perícias, gostaria que me contasse um pouco de como foi a "
               "experiência com a obra aí no condomínio {predio}. Tem algum horário que posso te ligar? "
               "Ou prefere que a gente converse por aqui?")


def esc(v):
    return html.escape(str(v).strip(), quote=True) if v is not None else ""


def texto(v):
    """Célula -> string limpa (sem espaços nas pontas); números viram texto sem '.0'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def norm(v):
    v = unicodedata.normalize("NFKD", texto(v).lower()).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", v).strip()


def so_digitos(v):
    return re.sub(r"\D", "", texto(v))


def id_youtube(url):
    """Aceita watch?v=, youtu.be/, shorts/, embed/, live/ ou só o ID."""
    u = texto(url)
    if not u:
        return ""
    m = re.search(r"(?:v=|youtu\.be/|shorts/|embed/|live/)([A-Za-z0-9_-]{11})", u)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{11}", u):
        return u
    return ""


def slug(v):
    return re.sub(r"[^a-z0-9]+", "-", norm(v)).strip("-")


def cidade_nome(v):
    """'Vila Velha – ES' -> 'Vila Velha'"""
    return re.split(r"\s*[–\-/]\s*ES\b", texto(v))[0].strip() or texto(v)


def link_http(url):
    u = texto(url)
    if u and not re.match(r"https?://", u):
        u = "https://" + u
    return u


# ---------------------------------------------------------------- leitura da planilha
def ler_planilha(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb["Obras"] if "Obras" in wb.sheetnames else wb.active

    # localiza a linha de cabeçalho (a que contém "Nº Obra")
    cab = None
    for r in range(1, 15):
        vals = [norm(c.value) for c in ws[r]]
        if any(v.startswith("nº obra") or v.startswith("n obra") or v.startswith("no obra") for v in vals):
            cab = r
            break
    if cab is None:
        sys.exit("Não encontrei a linha de cabeçalho (coluna 'Nº Obra') na planilha.")

    col = {}
    for c in ws[cab]:
        k = norm(c.value).replace("*", "").strip()
        if k:
            col[k] = c.column

    def coluna(*nomes):
        for n in nomes:
            for k, v in col.items():
                if k.startswith(n):
                    return v
        return None

    C = {
        "num": coluna("nº obra", "n obra", "no obra", "numero"),
        "predio": coluna("nome do predio", "nome do prédio", "predio"),
        "bairro": coluna("bairro"),
        "cidade": coluna("cidade"),
        "ano": coluna("ano"),
        "tipo": coluna("tipo de servico", "tipo de serviço"),
        "estagio": coluna("estagio", "estágio"),
        "categoria": coluna("tipo de obra", "categoria"),
        "titulo": coluna("titulo", "título"),
        "resumo": coluna("resumo"),
        "foto": coluna("foto"),
        "glink": coluna("link avaliacao", "link avaliação"),
        "av_nome": coluna("nome do avaliador"),
        "av_texto": coluna("texto da avaliacao", "texto da avaliação"),
        "video": coluna("link video", "link vídeo", "video", "vídeo"),
        "sindico": coluna("nome do sindico", "nome do síndico", "sindico"),
        "zap": coluna("whatsapp"),
    }
    faltando = [k for k in ("num", "predio", "tipo", "estagio", "titulo") if C[k] is None]
    if faltando:
        sys.exit(f"Colunas obrigatórias não encontradas na planilha: {faltando}")

    def cel(row, k):
        return row[C[k] - 1].value if C[k] else None

    obras = []
    for row in ws.iter_rows(min_row=cab + 1):
        num = texto(cel(row, "num"))
        predio = texto(cel(row, "predio"))
        if not num or not predio or not re.match(r"\d", num):
            continue
        obras.append({k: texto(cel(row, k)) for k in C})

    # aba "Página": pares rótulo | valor
    config = {}
    if "Página" in wb.sheetnames:
        for row in wb["Página"].iter_rows(min_row=1, max_col=2, values_only=True):
            if row[0] and row[1]:
                config[norm(row[0])] = texto(row[1])
    return obras, config


# ---------------------------------------------------------------- montagem do HTML
def agrupar(obras):
    """Agrupa por prédio preservando a ordem de aparição; mídia vem da 1ª linha preenchida."""
    grupos, ordem = {}, []
    for o in obras:
        chave = norm(o["predio"])
        if chave not in grupos:
            grupos[chave] = {"predio": o["predio"], "bairro": o["bairro"], "cidade": o["cidade"],
                             "foto": "", "glink": "", "av_nome": "", "av_texto": "", "video": "",
                             "sindico": "", "zap": "", "obras": []}
            ordem.append(chave)
        g = grupos[chave]
        for k in ("foto", "glink", "av_nome", "av_texto", "video", "bairro", "cidade"):
            if not g[k] and o[k]:
                g[k] = o[k]
        if not g["sindico"] and o["sindico"] and so_digitos(o["zap"]):
            g["sindico"], g["zap"] = o["sindico"], so_digitos(o["zap"])
        g["obras"].append(o)
    return [grupos[k] for k in ordem]


def html_obra(o, avisos, cidade):
    tipo = TIPOS.get(norm(o["tipo"]))
    if not tipo:
        avisos.append(f"Obra {o['num']}: tipo de serviço desconhecido '{o['tipo']}' (usando Gestão)")
        tipo = TIPOS["gestao"]
    est = next((e for e in ESTAGIOS if e[1] == norm(o["estagio"])), None)
    if not est:
        avisos.append(f"Obra {o['num']}: estágio desconhecido '{o['estagio']}' (usando Concluído)")
        est = ESTAGIOS[2]
    num = o["num"].zfill(3) if o["num"].isdigit() else o["num"]
    resumo = f'\n          <p class="obra-item-resumo">{esc(o["resumo"])}</p>' if o["resumo"] else ""
    ano = f'\n          <div class="obra-item-ano">{SVG_CAL}{esc(o["ano"])}</div>' if o["ano"] else ""
    return (
        f'        <div class="obra-item-block" data-servico="{tipo[0]}" data-estagio="{est[0]}" '
        f'data-obra="{slug(o["categoria"]) or "outros"}" data-cidade="{slug(cidade_nome(cidade))}">\n'
        f'          <div class="obra-item-meta-row">\n'
        f'            <span class="obra-item-num">Obra {esc(num)}</span>\n'
        f'            {tipo[1]}<span class="estagio-pill estagio-tag-{est[0]}">{est[2]}</span>\n'
        f'          </div>\n'
        f'          <div class="obra-item-titulo">{esc(o["titulo"])}</div>{resumo}{ano}\n'
        f'        </div>'
    )


def html_sindico(g):
    if not (g["sindico"] and g["zap"]):
        return ""
    zap = g["zap"]
    if len(zap) in (10, 11):
        zap = "55" + zap
    msg = quote(MSG_SINDICO.format(nome=g["sindico"], predio=g["predio"]), safe="")
    return (f'<a class="sindico-btn-card" href="https://wa.me/{zap}?text={msg}" target="_blank" rel="noopener">'
            f'{SVG_ZAP}Falar com o Síndico</a>')


def html_card(g, avisos):
    predio = esc(g["predio"])
    endereco = ", ".join(x for x in (texto(g["bairro"]), texto(g["cidade"])) if x)

    # foto
    if g["foto"]:
        arquivo = os.path.basename(g["foto"])
        if not os.path.exists(os.path.join(PASTA_FOTOS, arquivo)):
            avisos.append(f"{g['predio']}: foto '{arquivo}' não existe em public/images/obras/")
        foto = (f'      <img src="{URL_FOTOS}{esc(arquivo)}" alt="Obra no condomínio {predio}" loading="lazy">\n'
                f'      <div class="obra-foto-overlay"></div>')
    else:
        foto = f'      <div class="obra-foto-placeholder">{SVG_FOTO}<span>Foto em breve</span></div>'

    # avaliação google
    glink = link_http(g["glink"])
    avaliacao = ""
    if glink:
        nome_av = f'\n      <div class="avaliador-nome">{esc(g["av_nome"])}</div>' if g["av_nome"] else ""
        texto_av = (f'\n      <div class="divider-thin"></div>\n      <p class="review-text">{esc(g["av_texto"])}</p>'
                    if g["av_texto"] else "")
        if not g["av_texto"]:
            avisos.append(f"{g['predio']}: tem link de avaliação mas o texto da avaliação está vazio")
        sind_av = html_sindico(g)
        avaliacao = (
            f'\n    <div class="obra-avaliacao">\n'
            f'      <div class="google-header">{SVG_GOOGLE}<span class="google-label">Avaliação Google</span></div>\n'
            f'      <div class="google-nota-row"><span class="google-nota">5,0</span><span class="google-nota-max">/ 5</span></div>\n'
            f'      <div class="stars">{SVG_STAR * 5}</div>{nome_av}{texto_av}\n'
            f'      <div class="av-footer">\n'
            f'        <a class="google-link" href="{esc(glink)}" target="_blank" rel="noopener">Ver no Google {SVG_EXT}</a>\n'
            f'        {sind_av}\n'
            f'      </div>\n'
            f'    </div>'
        )
        footer_meta = ""
    else:
        footer_meta = html_sindico(g)

    # vídeo
    vid = id_youtube(g["video"])
    if g["video"] and not vid:
        avisos.append(f"{g['predio']}: não reconheci o link do YouTube '{g['video']}'")
    video = ""
    if vid:
        video = (
            f'\n    <div class="obra-video">\n'
            f'      <div class="video-placeholder" onclick="loadVideo(this,\'{vid}\')" role="button" '
            f'aria-label="Assistir depoimento em vídeo">\n'
            f'        <div class="play-ring">{SVG_PLAY}</div><span class="play-label">Depoimento em vídeo</span>\n'
            f'      </div>\n'
            f'    </div>'
        )

    classes = "obra-card" + ("" if video else " sem-video") + ("" if avaliacao else " sem-avaliacao-col")
    itens = "\n".join(html_obra(o, avisos, g["cidade"]) for o in g["obras"])
    return (
        f'  <div class="{classes}">\n'
        f'    <div class="obra-nome-mobile">{predio}<div class="obra-endereco-mobile">{SVG_PIN}{esc(endereco)}</div></div>\n'
        f'    <div class="obra-foto">\n{foto}\n    </div>\n'
        f'    <div class="obra-dados">\n'
        f'      <div class="obra-dados-top">\n'
        f'        <h2 class="obra-nome">{predio}</h2>\n'
        f'        <div class="obra-endereco">{SVG_PIN}{esc(endereco)}</div>\n'
        f'        <div class="obras-lista">\n{itens}</div>\n'
        f'      </div>\n'
        f'      <div class="obra-footer-meta">{footer_meta}</div>\n'
        f'    </div>{avaliacao}{video}\n'
        f'  </div>'
    )


def _grupo(label, grupo, opcoes, wide=False):
    """opcoes: lista de (slug, rótulo). Só gera o grupo se houver mais de uma opção."""
    if len(opcoes) < 2:
        return ""
    b = [f'  <div class="filter-group{" wide" if wide else ""}"><span class="filter-label">{label}:</span>',
         f'    <button class="filter-btn active" data-grupo="{grupo}" onclick="filtrar(\'todos\',this,\'{grupo}\')">Todos</button>']
    for s, r in opcoes:
        b.append(f'    <button class="filter-btn" data-grupo="{grupo}" onclick="filtrar(\'{s}\',this,\'{grupo}\')">{esc(r)}</button>')
    b.append("  </div>")
    return "\n".join(b)


def html_filtros(obras, grupos):
    serv = [(v[0], r) for k, v, r in ((k, v, {"gestao": "Gestão", "fiscalizacao": "Fiscalização",
            "gef": "Gestão e Fiscalização"}[v[0]]) for k, v in TIPOS.items())
            if k in {norm(o["tipo"]) for o in obras}]
    est_usados = {norm(o["estagio"]) for o in obras}
    est = [(s, r) for s, chave, r in ESTAGIOS if chave in est_usados]
    # tipo de obra: ordem por quantidade (mais comuns primeiro)
    cont = {}
    for o in obras:
        c = o["categoria"] or "Outros"
        cont[c] = cont.get(c, 0) + 1
    cats = sorted(cont, key=lambda c: (c == "Outros", -cont[c], c))
    cat = [(slug(c), c) for c in cats]
    # cidade: ordem por quantidade de obras
    ccont = {}
    for o in obras:
        n = cidade_nome(o["cidade"])
        if n:
            ccont[n] = ccont.get(n, 0) + 1
    cid = [(slug(n), n) for n in sorted(ccont, key=lambda n: (-ccont[n], n))]
    # linha 1: serviço e estágio lado a lado; linhas seguintes: tipo de obra e cidade, cada uma inteira
    linha1 = [x for x in (_grupo("Serviço", "servico", serv), _grupo("Estágio", "estagio", est)) if x]
    outras = [x for x in (_grupo("Tipo de obra", "obra", cat, wide=True),
                          _grupo("Cidade", "cidade", cid, wide=True)) if x]
    partes = ['\n  <span class="filter-sep"></span>\n'.join(linha1)] + outras
    return '<div class="filters-bar">\n' + "\n".join(p for p in partes if p) + '\n</div>'


def html_numeros(obras, grupos):
    cidades = {cidade_nome(o["cidade"]) for o in obras if o["cidade"]}
    avals = sum(1 for g in grupos if g["glink"])
    itens = [(len(obras), "obras acompanhadas"), (len(grupos), "condomínios atendidos"),
             (len(cidades), "cidades no Espírito Santo")]
    if avals:
        itens.append(("5,0", f"nota no Google · {avals} avaliações"))
    return "\n".join(f'  <div class="stat"><span class="stat-num">{n}</span><span class="stat-label">{l}</span></div>'
                      for n, l in itens)


def html_video_apresentacao(config, avisos):
    link = config.get("video de apresentacao (youtube)") or next(
        (v for k, v in config.items() if k.startswith("video de apresentacao")), "")
    if not link:
        return "", " sem-video"
    vid = id_youtube(link)
    if not vid:
        avisos.append(f"Página: não reconheci o link do vídeo de apresentação '{link}'")
        return "", " sem-video"
    return (f'  <div class="video-apresentacao">\n'
            f'    <div class="video-placeholder" style="background-image:url(https://i.ytimg.com/vi/{vid}/hqdefault.jpg)" '
            f'onclick="loadVideo(this,\'{vid}\')" role="button" aria-label="Assistir vídeo de apresentação">\n'
            f'      <div class="play-ring">{SVG_PLAY}</div><span class="play-label">Assista: como funciona</span>\n'
            f'    </div>\n  </div>'), ""


def main():
    ap = argparse.ArgumentParser(description="Gera public/obra.html a partir da planilha de obras.")
    ap.add_argument("--planilha", default=PLANILHA)
    ap.add_argument("--saida", default=SAIDA)
    args = ap.parse_args()

    obras, config = ler_planilha(args.planilha)
    if not obras:
        sys.exit("Nenhuma obra encontrada na planilha.")
    avisos = []
    grupos = agrupar(obras)
    cards = "\n".join(html_card(g, avisos) for g in grupos)

    with open(TEMPLATE, encoding="utf-8") as f:
        tpl = f.read()
    video_html, classe = html_video_apresentacao(config, avisos)
    saida = (tpl.replace("{{CARDS}}", cards)
             .replace("{{FILTROS}}", html_filtros(obras, grupos))
             .replace("{{NUMEROS}}", html_numeros(obras, grupos))
             .replace("{{VIDEO_APRESENTACAO}}", video_html)
             .replace("{{COMO_ATUA_CLASSE}}", classe))
    for o in obras:
        if not o["categoria"]:
            avisos.append(f"Obra {o['num']}: sem 'Tipo de Obra' (entra no filtro como Outros)")
    with open(args.saida, "w", encoding="utf-8") as f:
        f.write(saida)

    print(f"OK: {len(obras)} obras em {len(grupos)} prédios -> {os.path.relpath(args.saida, RAIZ)} "
          f"({os.path.getsize(args.saida) // 1024} KB)")
    print(f"    com foto: {sum(1 for g in grupos if g['foto'])} | com avaliação Google: "
          f"{sum(1 for g in grupos if g['glink'])} | com vídeo: {sum(1 for g in grupos if id_youtube(g['video']))} | "
          f"com botão do síndico: {sum(1 for g in grupos if g['sindico'] and g['zap'])}")
    for a in avisos:
        print("AVISO:", a)


if __name__ == "__main__":
    main()
